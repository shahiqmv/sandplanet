"""Transactions from a spreadsheet (FINANCE_BUILD_BRIEF.md, stage 2) — how
January to June 2026, before PLANET held the operations, is brought into the
books without keying each form.

One row is one transaction with one line; a row with no Type adds a line to
the transaction above it. Each transaction goes through the same rules as
its form (core/books.py), so nothing gets in this way that the form would
refuse. The file is checked first and nothing is saved; it is then imported
whole or not at all. The same file is not taken twice, and a batch can be
undone.
"""
import hashlib
from decimal import Decimal

from django.db import transaction

from . import books
from .audit import audit
from .ledger import BASE, ZERO, q2
from .models import LedgerAccount, LedgerImport, LedgerTxn, Site
from .reconcile import _cell_date, _cell_money

MAX_TXNS = 2000
TYPES = {"expense": "EXPENSE", "deposit": "DEPOSIT", "transfer": "TRANSFER",
         "bill": "BILL", "invoice": "INVOICE", "sales invoice": "INVOICE"}
GST_WORDS = {"": "NONE", "none": "NONE", "no gst": "NONE", "no": "NONE",
             "standard": "STANDARD", "standard rated": "STANDARD",
             "gst": "STANDARD", "zero": "ZERO", "zero rated": "ZERO",
             "exempt": "EXEMPT", "out of scope": "OUT_OF_SCOPE"}
# (key, heading in the sheet, column width, what goes in it)
COLUMNS = [
    ("type", "Type", 12,
     "Expense, Deposit, Transfer, Bill or Invoice. Leave blank to add "
     "another line to the transaction above."),
    ("date", "Date", 12, "The date of the transaction (of the bill or "
     "invoice)."),
    ("account", "Bank / cash account", 24,
     "Expense: paid from. Deposit: paid into. Transfer: from. Its name or "
     "code as in the chart. Blank for a bill or invoice."),
    ("to_account", "To account", 22, "Transfer only: the account the money "
     "went to."),
    ("party", "Name", 30, "Who was paid, who paid, the supplier or the "
     "customer."),
    ("reference", "Reference", 18,
     "Bill: the supplier's bill number. Invoice: the invoice number as "
     "issued. Otherwise the cheque, transfer or slip reference."),
    ("line_account", "Account", 30,
     "What it was for: the account's code, or code and name (6320 or "
     "6320 Electricity). Not for a transfer."),
    ("description", "Description", 32, "What this line is."),
    ("amount", "Amount", 14, "The line's amount. For a transfer, the amount "
     "sent."),
    ("gst", "GST", 14, "Blank / None, Standard, Zero rated, Exempt or Out "
     "of scope."),
    ("gst_incl", "Amount includes GST", 12,
     "Y if the amount is the total including GST; blank if it is before "
     "GST."),
    ("tax_invoice", "Tax invoice held", 10,
     "Y if we hold the supplier's valid tax invoice (expense, bill). The "
     "GST is only claimed when Y."),
    ("tax_invoice_no", "Tax invoice no.", 16, "Expense: the supplier's tax "
     "invoice number."),
    ("tin", "TIN", 16, "The supplier's or customer's TIN."),
    ("due_date", "Due date", 12, "Bill, invoice: when it falls due. Blank "
     "uses their credit days."),
    ("currency", "Currency", 9, "Bill, invoice: blank for MVR, or USD etc."),
    ("rate", "Rate to MVR", 11, "For anything not in rufiyaa: the rate on "
     "the day."),
    ("amount_to", "Amount received", 14, "Transfer between currencies: what "
     "arrived, in the other account's currency."),
    ("control", "Payable / receivable account", 24,
     "Bill, invoice: only if not the usual one (2010 / 1210)."),
    ("site", "Site", 8, "The site code, if the line belongs to a site."),
    ("memo", "Memo", 30, "A note on the whole transaction."),
]
KEYS = {head.lower(): key for key, head, _, _ in COLUMNS}


def template():
    """The empty sheet to fill, with how to fill it and what to pick from."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    for i, (_, head, width, _) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=head)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="23272F")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    how = wb.create_sheet("How to fill")
    how.column_dimensions["A"].width = 30
    how.column_dimensions["B"].width = 110
    rows = [("One row is one transaction.", "A row with Type left blank is "
             "another line of the transaction above it (a bill with two "
             "kinds of cost, say)."),
            ("Nothing is saved until it is all right.", "The file is checked "
             "first and every problem is listed by row. It is then imported "
             "whole, or not at all."),
            ("", "")]
    rows += [(head, note) for _, head, _, note in COLUMNS]
    rows += [("", ""), ("Examples", ""),
             ("Expense", "Type Expense · Date 15/01/2026 · Bank BML MVR "
              "Current · Name STELCO · Account 6320 · Amount 1080 · GST "
              "Standard · Amount includes GST Y · Tax invoice held Y · Tax "
              "invoice no. ST-551 · TIN 1000123GST501"),
             ("Bill", "Type Bill · Date 20/01/2026 · Name Manas Hardware · "
              "Reference MH-1001 · Account 5110 · Amount 12000 · GST "
              "Standard · Tax invoice held Y · TIN 1012345GST501"),
             ("Transfer", "Type Transfer · Date 03/02/2026 · Bank BML MVR "
              "Current · To account Bank of Maldives USD · Amount 15420 · "
              "Amount received 1000")]
    for r, (a, b) in enumerate(rows, 1):
        how.cell(row=r, column=1, value=a).font = Font(bold=True)
        how.cell(row=r, column=2, value=b).alignment = Alignment(wrap_text=True)

    acc = wb.create_sheet("Accounts")
    acc.column_dimensions["A"].width = 10
    acc.column_dimensions["B"].width = 46
    acc.column_dimensions["C"].width = 26
    acc.column_dimensions["D"].width = 10
    for i, h in enumerate(["Code", "Name", "Type", "Currency"], 1):
        acc.cell(row=1, column=i, value=h).font = Font(bold=True)
    for r, a in enumerate(LedgerAccount.objects.filter(
            is_active=True, is_group=False).order_by("code"), 2):
        for i, v in enumerate([a.code, a.name, a.get_type_display(),
                               a.currency or BASE], 1):
            acc.cell(row=r, column=i, value=v)
    sites = wb.create_sheet("Sites")
    sites.cell(row=1, column=1, value="Code").font = Font(bold=True)
    sites.cell(row=1, column=2, value="Name").font = Font(bold=True)
    for r, s in enumerate(Site.objects.order_by("code"), 2):
        sites.cell(row=r, column=1, value=s.code)
        sites.cell(row=r, column=2, value=s.name)
    return wb


def _yes(v):
    return str(v or "").strip().lower() in ("y", "yes", "true", "1", "x")


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)                    # a cheque number Excel made a float
    return " ".join(str(v).split())


class _Lookup:
    """Accounts and sites by what a person would type for them."""

    def __init__(self):
        self.accounts = {}
        for a in LedgerAccount.objects.all():
            for k in (a.code, a.name, f"{a.code} {a.name}"):
                self.accounts.setdefault(k.strip().lower(), a)
        self.sites = {s.code.lower(): s for s in Site.objects.all()}

    def account(self, v, what):
        key = _text(v).lower()
        if not key:
            return None
        a = self.accounts.get(key) or self.accounts.get(key.split(" ")[0])
        if a is None:
            raise ValueError(f"{what} “{_text(v)}” is not in the chart of "
                             "accounts.")
        return a


def _read(upload):
    """The sheet as [(row number, {key: cell})]. Raises ValueError."""
    from openpyxl import load_workbook
    name = (getattr(upload, "name", "") or "").lower()
    if not name.endswith((".xlsx", ".xlsm")):
        raise ValueError("Import the Excel template (.xlsx).")
    try:
        wb = load_workbook(upload, read_only=True, data_only=True)
    except Exception:
        raise ValueError("That Excel file can't be opened.")
    ws = wb["Transactions"] if "Transactions" in wb.sheetnames else wb.active
    rows = ws.iter_rows(values_only=True)
    head = next(rows, None) or []
    cols = {i: KEYS[str(h).strip().lower()] for i, h in enumerate(head)
            if h is not None and str(h).strip().lower() in KEYS}
    missing = [h for k, h, _, _ in COLUMNS
               if k in ("type", "date", "amount") and k not in cols.values()]
    if missing:
        raise ValueError("This isn't the import template — the first row "
                         f"has no {', '.join(missing)} column. Download the "
                         "template and fill that.")
    out = []
    for n, row in enumerate(rows, 2):
        cells = {cols[i]: v for i, v in enumerate(row) if i in cols}
        if any(v not in (None, "") for v in cells.values()):
            out.append((n, cells))
    return out


def _group(rows):
    """Rows into transactions: a row with a Type starts one."""
    txns = []
    for n, c in rows:
        if _text(c.get("type")):
            txns.append({"row": n, "head": c, "rows": [(n, c)]})
        elif not txns:
            raise ValueError(f"Row {n}: the first row needs a Type.")
        else:
            txns[-1]["rows"].append((n, c))
    return txns


def _line(n, c, look, gst_rate):
    amount = _cell_money(c.get("amount"))
    if amount is None:
        raise ValueError(f"Row {n}: the amount is missing or not a number.")
    word = _text(c.get("gst")).lower()
    if word not in GST_WORDS:
        raise ValueError(f"Row {n}: GST “{_text(c.get('gst'))}” — use None, "
                         "Standard, Zero rated, Exempt or Out of scope.")
    treat = GST_WORDS[word]
    gst = None
    if treat == "STANDARD":
        if _yes(c.get("gst_incl")):
            net = q2(amount / (1 + Decimal(gst_rate) / 100))
            amount, gst = net, amount - net
    acc = look.account(c.get("line_account"), f"Row {n}: account")
    if acc is None:
        raise ValueError(f"Row {n}: say which account it is for.")
    site = None
    if _text(c.get("site")):
        site = look.sites.get(_text(c.get("site")).lower())
        if site is None:
            raise ValueError(f"Row {n}: there is no site "
                             f"“{_text(c.get('site'))}”.")
    return {"account": acc.id, "description": _text(c.get("description")),
            "amount": str(amount), "gst_treatment": treat,
            "gst_amount": None if gst is None else str(gst),
            "site": site.id if site else None}


def _data(t, look, gst_rate):
    """What the form would have sent for this transaction."""
    n, h = t["row"], t["head"]
    typ = TYPES.get(_text(h.get("type")).lower())
    if typ is None:
        raise ValueError(f"Row {n}: Type “{_text(h.get('type'))}” — use "
                         "Expense, Deposit, Transfer, Bill or Invoice.")
    d = _cell_date(h.get("date"))
    if d is None:
        raise ValueError(f"Row {n}: the date is missing or can't be read.")
    rate = _cell_money(h.get("rate"))
    data = {"type": typ, "date": d.isoformat(),
            "party": _text(h.get("party")),
            "party_tin": _text(h.get("tin")),
            "reference": _text(h.get("reference")),
            "memo": _text(h.get("memo")),
            "fx_rate": None if rate is None else str(h.get("rate")).strip(),
            "tax_invoice_held": _yes(h.get("tax_invoice")),
            "tax_invoice_no": _text(h.get("tax_invoice_no"))}
    if typ in books.DOC_TYPES:
        ctrl = look.account(h.get("control"), f"Row {n}: account")
        data.update(account=ctrl.id if ctrl else None,
                    currency=_text(h.get("currency")).upper() or BASE)
        due = _cell_date(h.get("due_date"))
        if _text(h.get("due_date")) and due is None:
            raise ValueError(f"Row {n}: the due date can't be read.")
        data["due_date"] = due.isoformat() if due else None
    else:
        bank = look.account(h.get("account"), f"Row {n}: bank / cash account")
        if bank is None:
            raise ValueError(f"Row {n}: say which bank or cash account.")
        data["account"] = bank.id
    if typ == "TRANSFER":
        to = look.account(h.get("to_account"), f"Row {n}: to account")
        if to is None:
            raise ValueError(f"Row {n}: say which account the money went to.")
        if len(t["rows"]) > 1:
            raise ValueError(f"Row {t['rows'][1][0]}: a transfer has one "
                             "row.")
        amount = _cell_money(h.get("amount"))
        if amount is None:
            raise ValueError(f"Row {n}: the amount is missing.")
        arrive = _cell_money(h.get("amount_to"))
        data.update(to_account=to.id, amount=str(amount),
                    amount_to=None if arrive is None else str(arrive))
    else:
        data["lines"] = [_line(rn, c, look, gst_rate) for rn, c in t["rows"]]
    return typ, data


class _Undo(Exception):
    pass


def run(upload, actor, commit=False):
    """Check the file, or import it. Returns (result, error).

    The check runs the real thing and rolls it back, so what it reports is
    exactly what importing would do."""
    raw = upload.read()
    sha = hashlib.sha256(raw).hexdigest()
    upload.seek(0)
    before = LedgerImport.objects.filter(sha256=sha, undone=False).first()
    if before:
        return None, (f"This file was already imported on "
                      f"{before.created_at:%d %b %Y} "
                      f"({before.count} transactions). Importing it again "
                      "would enter everything twice.")
    try:
        txns = _group(_read(upload))
    except ValueError as exc:
        return None, str(exc)
    if not txns:
        return None, "There are no transactions in the file."
    if len(txns) > MAX_TXNS:
        return None, (f"That is {len(txns):,} transactions — import at most "
                      f"{MAX_TXNS:,} at a time (a month or two per file).")
    look, gst_rate = _Lookup(), books.gst_rate()
    rows, errors, total, batch, made = [], 0, ZERO, None, []
    try:
        with transaction.atomic():
            for t in txns:
                row = {"row": t["row"], "lines": len(t["rows"]),
                       "type": _text(t["head"].get("type")),
                       "party": _text(t["head"].get("party")),
                       "reference": _text(t["head"].get("reference")),
                       "date": None, "amount": None, "number": "",
                       "error": None}
                try:
                    typ, data = _data(t, look, gst_rate)
                    row["date"] = data["date"]
                    saved, msg = books.save_txn(typ, data, actor)
                    if msg:
                        raise ValueError(f"Row {t['row']}: {msg}")
                    made.append(saved.pk)
                    row.update(number=saved.number, currency=saved.currency,
                               amount=saved.amount,
                               amount_mvr=saved.amount_mvr)
                    total += saved.amount_mvr
                except ValueError as exc:
                    row["error"] = str(exc)
                    errors += 1
                rows.append(row)
            if errors or not commit:
                raise _Undo
            # the file is kept only once everything in it has gone in
            upload.seek(0)
            batch = LedgerImport.objects.create(
                file=upload, filename=(upload.name or "import.xlsx")[:200],
                sha256=sha, created_by=actor, count=len(rows),
                total_mvr=total)
            LedgerTxn.objects.filter(pk__in=made).update(import_batch=batch)
    except _Undo:
        batch = None
        for r in rows:
            r["number"] = ""          # the numbers were handed back
    if batch:
        audit("ledger_import", batch.id, "IMPORTED", actor=actor,
              detail={"file": batch.filename, "count": batch.count,
                      "total_mvr": str(total)})
    return {"rows": rows, "count": len(rows), "errors": errors,
            "total_mvr": total, "imported": bool(batch),
            "batch": batch.id if batch else None}, None


def undo(batch, actor, reason):
    """Void every transaction of a batch — all, or none."""
    reason = (reason or "").strip()
    if batch.undone:
        return "This import has already been undone."
    if not reason:
        return "Say why it is being undone."
    try:
        with transaction.atomic():
            # payments first, so the bills they settle are free to void
            order = sorted(batch.txns.exclude(status="VOID"),
                           key=lambda t: t.type not in books.PAY_TYPES)
            for t in order:
                msg = books.void_txn(t, actor,
                                     f"Import undone — {reason}"[:300])
                if msg:
                    raise ValueError(f"{t.number}: {msg}")
            batch.undone = True
            batch.save(update_fields=["undone"])
    except ValueError as exc:
        return (f"The import can't be undone as a whole — {exc} Nothing was "
                "changed.")
    audit("ledger_import", batch.id, "IMPORT_UNDONE", actor=actor,
          detail={"file": batch.filename, "reason": reason[:200]})
    return None


def history():
    return [{"id": b.id, "filename": b.filename, "count": b.count,
             "total_mvr": b.total_mvr, "undone": b.undone,
             "file_url": b.file.url if b.file else None,
             "by": b.created_by.full_name if b.created_by_id else "",
             "at": b.created_at} for b in
            LedgerImport.objects.select_related("created_by")[:100]]

