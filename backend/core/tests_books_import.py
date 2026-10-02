"""Transactions from a spreadsheet: checked first, imported whole or not at
all, never twice, and undone as a batch (FINANCE_BUILD_BRIEF.md, stage 2)."""
import io
from datetime import datetime
from decimal import Decimal as D

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from . import books_import
from .models import AuditLog, LedgerImport, LedgerParty, LedgerTxn, Site
from .tests_books import BooksBase

HEAD = [h for _, h, _, _ in books_import.COLUMNS]


def sheet(rows, head=HEAD, name="jan.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(head)
    for r in rows:
        ws.append([r.get(h) for h in head])
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(name, buf.getvalue())


GOOD = [
    {"Type": "Expense", "Date": datetime(2026, 1, 15),
     "Bank / cash account": "BML MVR", "Name": "STELCO",
     "Reference": 4471.0, "Account": "6320 Electricity", "Amount": 1080,
     "GST": "Standard", "Amount includes GST": "Y", "Tax invoice held": "Y",
     "Tax invoice no.": "ST-551", "TIN": "1000123GST501"},
    {"Type": "Bill", "Date": "20/01/2026", "Name": "Manas Hardware",
     "Reference": "MH-1001", "Account": "5110", "Description": "Cement",
     "Amount": "12,000.00", "GST": "standard"},
    # a second line of the same bill
    {"Account": "6320", "Description": "Delivery", "Amount": 500},
    {"Type": "Deposit", "Date": "2026-01-22", "Bank / cash account": "BML MVR",
     "Name": "Soneva Fushi", "Account": "4120", "Amount": 54000},
    {"Type": "Transfer", "Date": "03-02-2026", "Bank / cash account": "BML MVR",
     "To account": "BML USD", "Amount": 15420, "Amount received": 1000},
    {"Type": "Invoice", "Date": "05/02/2026", "Name": "Kuramathi",
     "Reference": "INV-2026-0003", "Account": "4120", "Amount": 2000,
     "GST": "Zero rated", "Currency": "usd", "Rate to MVR": 15.42,
     "Due date": "05/03/2026"},
]


class ImportTests(BooksBase):
    def post(self, f, commit=False):
        return self.c.post("/api/v1/ledger/imports",
                           {"file": f, "commit": "1" if commit else "0"},
                           format="multipart")

    def test_checking_saves_nothing_and_importing_saves_it_all(self):
        before = (LedgerTxn.objects.count(), AuditLog.objects.count())
        r = self.post(sheet(GOOD))
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual((r.data["count"], r.data["errors"],
                          r.data["imported"]), (5, 0, False))
        self.assertEqual([x["lines"] for x in r.data["rows"]], [1, 2, 1, 1, 1])
        self.assertEqual((LedgerTxn.objects.count(), AuditLog.objects.count(),
                          LedgerParty.objects.count()), (*before, 0))
        # one file, byte for byte: a workbook built again a second later
        # carries a different timestamp inside and is a different file
        raw = sheet(GOOD).read()
        r = self.post(SimpleUploadedFile("jan.xlsx", raw), commit=True)
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual([x["number"] for x in r.data["rows"]],
                         ["EXP-001", "BILL-001", "DEP-001", "TRF-001",
                          "SALE-001"])
        # each went through its form's rules
        exp = LedgerTxn.objects.get(number="EXP-001")
        self.assertEqual((exp.amount, exp.reference, exp.tax_invoice_held,
                          exp.tax_invoice_no, str(exp.date)),
                         (D("1080.00"), "4471", True, "ST-551", "2026-01-15"))
        self.assertEqual(self.bal("1430"), D("80.00"))       # GST taken out
        bill = LedgerTxn.objects.get(number="BILL-001")
        self.assertEqual((bill.lines.count(), bill.amount),
                         (2, D("13460.00")))                 # 12,960 + 500
        self.assertEqual(self.bal("2010"), D("-13460.00"))
        inv = LedgerTxn.objects.get(number="SALE-001")
        self.assertEqual((inv.currency, inv.amount_mvr, str(inv.due_date)),
                         ("USD", D("30840.00"), "2026-03-05"))
        self.assertEqual(self.bal(self.usd), D("15420.00"))
        b = LedgerImport.objects.get()
        self.assertEqual((b.count, b.txns.count(), b.filename),
                         (5, 5, "jan.xlsx"))
        # the same file is not taken twice
        self.assertIn("already imported", self.post(
            SimpleUploadedFile("jan-copy.xlsx", raw)).data["detail"])

    def test_one_bad_row_and_nothing_goes_in(self):
        rows = [dict(GOOD[0]), dict(GOOD[3], Account="9999 Nowhere"),
                dict(GOOD[1], Date="not a date"),
                dict(GOOD[4], Type="Journal"),
                dict(GOOD[0], GST="Reduced"),
                dict(GOOD[0], Date=datetime(2025, 6, 1))]
        r = self.post(sheet(rows), commit=True)
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual((r.data["errors"], r.data["imported"]), (5, False))
        errs = [x["error"] or "" for x in r.data["rows"]]
        self.assertIsNone(r.data["rows"][0]["error"])
        self.assertIn("Row 3: account “9999 Nowhere” is not in the chart",
                      errs[1])
        self.assertIn("Row 4: the date", errs[2])
        self.assertIn("Row 5: Type “Journal”", errs[3])
        self.assertIn("Row 6: GST “Reduced”", errs[4])
        self.assertIn("Row 7: The books start on", errs[5])
        self.assertEqual((LedgerTxn.objects.count(),
                          LedgerImport.objects.count()), (0, 0))
        # the good row's number was handed back
        self.assertEqual(r.data["rows"][0]["number"], "")
        self.assertEqual(self.txn("EXPENSE", party="X", lines=[
            {"account": self.acc["6320"].id, "amount": "5"}]).data["number"],
            "EXP-001")

    def test_the_wrong_file_says_so(self):
        self.assertIn("isn't the import template", self.post(
            sheet([{"a": 1}], head=["a", "b"])).data["detail"])
        self.assertIn("Import the Excel template", self.post(
            SimpleUploadedFile("x.csv", b"Type,Date\n")).data["detail"])
        self.assertIn("no transactions", self.post(sheet([])).data["detail"])
        self.assertIn("first row needs a Type", self.post(
            sheet([{"Account": "6320", "Amount": 5}])).data["detail"])
        self.assertIn("Attach", self.c.post(
            "/api/v1/ledger/imports", {}, format="multipart").data["detail"])

    def test_a_site_and_a_same_bill_twice(self):
        Site.objects.create(code="SJR", name="SJR")
        rows = [dict(GOOD[1], Site="sjr"), dict(GOOD[1])]
        r = self.post(sheet(rows))
        self.assertIsNone(r.data["rows"][0]["error"])
        self.assertIn("already entered", r.data["rows"][1]["error"])
        self.assertIn("no site", self.post(
            sheet([dict(GOOD[1], Site="ZZZ9")])).data["rows"][0]["error"])

    def test_a_batch_is_undone_whole_or_not_at_all(self):
        b = self.post(sheet(GOOD), commit=True).data["batch"]
        url = f"/api/v1/ledger/imports/{b}/undo"
        self.assertIn("Say why", self.c.post(url, {}, format="json")
                      .data["detail"])
        # the bill has since been paid by hand: the batch can't just vanish
        bill = LedgerTxn.objects.get(number="BILL-001")
        pay = self.c.post("/api/v1/ledger/txns", {
            "type": "BILL_PAY", "date": "2026-02-10", "account": self.bank.id,
            "applies": [{"doc": bill.id, "amount": "13460"}]}, format="json")
        self.assertEqual(pay.status_code, 201, pay.data)
        r = self.c.post(url, {"reason": "Wrong month"}, format="json")
        self.assertIn("can't be undone as a whole", r.data["detail"])
        self.assertIn("BILL-001", r.data["detail"])
        self.assertEqual(LedgerTxn.objects.filter(status="VOID").count(), 0)
        self.c.post(f"/api/v1/ledger/txns/{pay.data['id']}/void",
                    {"reason": "x"}, format="json")
        r = self.c.post(url, {"reason": "Wrong month"}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertTrue(r.data["imports"][0]["undone"])
        self.assertEqual(LedgerTxn.objects.filter(
            import_batch_id=b, status="POSTED").count(), 0)
        self.assertEqual((self.bal(self.bank), self.bal("2010"),
                          self.bal("1210")),
                         (D("1000000.00"), D("0.00"), D("0.00")))
        # undone, the file can be corrected and brought in again
        self.assertEqual(self.post(sheet(GOOD), commit=True).status_code, 201)

    def test_the_template_is_the_sheet_the_import_reads(self):
        r = self.c.get("/api/v1/ledger/import/template")
        self.assertIn("spreadsheet", r["Content-Type"])
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual([c.value for c in wb["Transactions"][1]], HEAD)
        self.assertIn("6320", [row[0].value for row in wb["Accounts"]])
        # read-only roles see the history and nothing more
        from rest_framework.test import APIClient
        sig = APIClient()
        sig.force_authenticate(self.sig)
        self.assertEqual(sig.get("/api/v1/ledger/imports").status_code, 200)
        self.assertEqual(sig.post("/api/v1/ledger/imports", {"file": sheet(GOOD)},
                                  format="multipart").status_code, 403)
