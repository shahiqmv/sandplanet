"""The GST return's figures and its two statements, out of the books
(FINANCE_BUILD_BRIEF.md — MIRA compliance).

Laid out on MIRA's own forms as published (mira.gov.mv, read 2026-10-02):
MIRA 205 "GST Return — General Goods and Services" v25.1, the Input Tax
Statement v25.1 and the Output Tax Statement v25.1. The box numbers and
column headings below are the forms'. If MIRA issues a new version, this is
the one file to bring up to date.

What goes on the return comes from the transaction forms (core/books.py):
a sale is an invoice or a deposit whose lines carry a GST treatment; input
tax is the GST on an expense or bill for which a valid tax invoice is held.
A line with no treatment is not a supply and is left off — a loan received,
the owner's money. Anything else on the GST accounts in the period (a
journal, the payment to MIRA) is listed apart, so the accountant sees it.
"""
from decimal import ROUND_HALF_UP, Decimal

from . import ledger
from .ledger import BASE, ZERO, q2
from .models import (CompanyParameter, JournalLine, LedgerTxn)

FORM = "MIRA 205 v25.1"
# the Input Tax Statement has a column for GST charged at each of these
INPUT_RATES = (6, 8, 12, 16, 17)
BOXES = [
    (1, "Sales of supplies subject to GST at {rate}% (inclusive of GST)"),
    (2, "Sales of zero-rated supplies"),
    (3, "Sales of exempt supplies"),
    (4, "Sales of supplies which are out of scope of GST"),
    (5, "Total sales (sum of Boxes 1 to 4)"),
    (6, "Output tax"),
    (7, "Input tax (attach the Statement of Input Tax)"),
    (8, "GST in respect of irrecoverable debts written off, and GST "
        "relating to credit notes spanning a rate change"),
    (9, "GST collected in excess"),
    (10, "GST liability for the period (Box 6 minus Box 7 and Box 8, plus "
         "Box 9)"),
]


def _param(key, default=""):
    try:
        v = CompanyParameter.objects.get(key=key).value
        return default if v in (None, "") else v
    except CompanyParameter.DoesNotExist:
        return default


def company():
    from .pdf import company_info
    info = company_info()
    return {"tin": info["tin"], "name": info["legal_name"],
            "activity_no": str(_param("gst_activity_no", ""))}


def _rufiyaa(v):
    """The return is filled in rounded to the nearest rufiyaa."""
    return v.quantize(Decimal("1"), rounding=ROUND_HALF_UP)


def _mvr(txn, amount):
    return q2(amount * (txn.fx_rate or 1)) if txn.currency != BASE else amount


def build(date_from, date_to):
    """Everything the GST page shows for a taxable period."""
    from .books import gst_rate
    co = company()
    base = (LedgerTxn.objects.filter(status="POSTED", is_opening=False,
                                     date__gte=date_from, date__lte=date_to)
            .select_related("party_ref").prefetch_related("lines__account")
            .order_by("date", "id"))
    out_rows, in_rows, warnings, used = [], [], [], set()
    tot = {k: ZERO for k in ("standard", "zero", "exempt", "oos",
                             "output_tax", "input_tax", "input_net")}
    untreated = 0
    for t in base.filter(type__in=("INVOICE", "DEPOSIT")):
        row = {"standard": ZERO, "zero": ZERO, "exempt": ZERO, "oos": ZERO,
               "gst": ZERO}
        for ln in t.lines.all():
            net, gst = _mvr(t, ln.amount), _mvr(t, ln.gst_amount)
            if ln.gst_treatment == "STANDARD":
                row["standard"] += net
                row["gst"] += gst
            elif ln.gst_treatment == "ZERO":
                row["zero"] += net
            elif ln.gst_treatment == "EXEMPT":
                row["exempt"] += net
            elif ln.gst_treatment == "OUT_OF_SCOPE":
                row["oos"] += net
            elif t.type == "INVOICE":
                untreated += 1
        if not any(row.values()):
            continue
        used.add(t.journal_id)
        out_rows.append({
            "txn": t.id, "type": t.type, "number": t.number,
            "customer_tin": t.party_tin or (t.party_ref.tin if t.party_ref
                                            else ""),
            "customer": t.party,
            "invoice_no": t.tax_invoice_no or t.reference or t.number,
            "date": t.date, "activity_no": co["activity_no"], **row})
        for k in ("standard", "zero", "exempt", "oos"):
            tot[k] += row[k]
        tot["output_tax"] += row["gst"]
    if untreated:
        warnings.append(
            f"{untreated} invoice line{'s have' if untreated > 1 else ' has'}"
            " no GST treatment and "
            f"{'are' if untreated > 1 else 'is'} left off the return. A sale "
            "is standard-rated, zero-rated, exempt or out of scope — open "
            "the invoice and choose.")

    old = 0
    for t in base.filter(type__in=("EXPENSE", "BILL"), tax_invoice_held=True):
        lines = list(t.lines.all())
        gst = sum((_mvr(t, ln.gst_amount) for ln in lines), ZERO)
        if not gst:
            continue
        used.add(t.journal_id)
        by_rate = {r: ZERO for r in INPUT_RATES}
        other = ZERO
        for ln in lines:
            g = _mvr(t, ln.gst_amount)
            if not g:
                continue
            r = int(ln.gst_rate) if ln.gst_rate == int(ln.gst_rate) else None
            if r in by_rate:
                by_rate[r] += g
            else:
                other += g
        capital = any(ln.account.type == "FIXED_ASSET" for ln in lines
                      if ln.gst_amount)
        inv_date = t.tax_invoice_date or t.date
        # MIRA: not claimable more than 12 months after the end of the
        # period it could first have been claimed in
        stale = (date_to - inv_date).days > 366 + 31
        old += stale
        net = sum((_mvr(t, ln.amount) for ln in lines), ZERO)
        in_rows.append({
            "txn": t.id, "type": t.type, "number": t.number,
            "supplier_tin": t.party_tin, "supplier": t.party,
            "invoice_no": t.tax_invoice_no or t.reference,
            "date": inv_date, "net": net,
            "gst": {str(r): v for r, v in by_rate.items()},
            "gst_other": other, "gst_total": gst,
            "activity_no": co["activity_no"],
            "kind": "Capital" if capital else "Revenue",
            "missing_tin": not t.party_tin,
            "missing_no": not (t.tax_invoice_no or t.reference),
            "stale": stale})
        tot["input_tax"] += gst
        tot["input_net"] += net
    missing = sum(1 for r in in_rows if r["missing_tin"] or r["missing_no"])
    if missing:
        warnings.append(
            f"{missing} purchase{'s' if missing > 1 else ''} on the input "
            "tax statement lack"
            f"{'' if missing > 1 else 's'} the supplier's TIN or the tax "
            "invoice number. MIRA's statement asks for both.")
    if old:
        warnings.append(
            f"{old} tax invoice{'s are' if old > 1 else ' is'} dated more "
            "than 12 months before this period ends — MIRA does not allow "
            "input tax that late.")
    if any(r["gst_other"] for r in in_rows):
        warnings.append("Some input tax was charged at a rate MIRA's "
                        "statement has no column for — check those lines.")
    if not co["activity_no"] and (out_rows or in_rows):
        warnings.append("The statements have a column for your taxable "
                        "activity number — it is not set yet (company "
                        "parameter gst_activity_no).")

    # anything else that touched the GST accounts in the period
    others = []
    keys = {"OUTPUT_GST": "GST payable", "INPUT_GST": "GST input tax"}
    for key, label in keys.items():
        acc = ledger.account_for(key)
        if acc is None:
            continue
        qs = (JournalLine.objects.filter(
            account=acc, entry__status="POSTED",
            entry__date__gte=date_from, entry__date__lte=date_to)
            .exclude(entry_id__in=[j for j in used if j])
            .select_related("entry", "entry__reversal_of",
                            "entry__reversed_by"))
        for ln in qs:
            e = ln.entry
            mirror = e.reversal_of if e.reversal_of_id else getattr(
                e, "reversed_by", None)
            if mirror is not None and mirror.date == e.date:
                continue                 # an edit or a void: cancels out
            others.append({"entry": e.id, "ref": e.ref, "date": e.date,
                           "account": label, "memo": e.memo,
                           "debit": ln.debit, "credit": ln.credit})

    rate = gst_rate()
    b = {1: tot["standard"] + tot["output_tax"], 2: tot["zero"],
         3: tot["exempt"], 4: tot["oos"], 6: tot["output_tax"],
         7: tot["input_tax"], 8: ZERO, 9: ZERO}
    r = {k: _rufiyaa(v) for k, v in b.items()}
    r[5] = r[1] + r[2] + r[3] + r[4]
    r[10] = r[6] - r[7] - r[8] + r[9]
    boxes = [{"box": n, "label": label.format(rate=f"{rate:g}"),
              "amount": r[n], "derived": n not in (8, 9)}
             for n, label in BOXES]
    return {
        "form": FORM, "company": co, "date_from": date_from,
        "date_to": date_to, "gst_rate": rate, "boxes": boxes,
        "exact": {"output_tax": tot["output_tax"],
                  "input_tax": tot["input_tax"],
                  "net": tot["output_tax"] - tot["input_tax"]},
        "output": out_rows, "input": in_rows,
        "output_total": {k: tot[k] for k in ("standard", "zero", "exempt",
                                             "oos", "output_tax")},
        "input_total": {"net": tot["input_net"], "gst": tot["input_tax"]},
        "other_entries": others, "warnings": warnings,
    }


def _book(sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows, widths, bold in sheets:
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
        for i in bold:
            for c in ws[i]:
                c.font = Font(bold=True)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows():
            for c in row:
                if hasattr(c.value, "quantize"):
                    c.number_format = "#,##0.00"
    return wb


def output_statement_xlsx(data):
    """MIRA's Output Tax Statement v25.1: a line per invoice, and the
    summary per taxable activity."""
    head = ["Customer TIN", "Customer Name", "Invoice No.", "Invoice Date",
            f"Value of Supplies Subject to GST at {data['gst_rate']:g}% or "
            "17% (excluding GST)", "Value of Zero-Rated Supplies",
            "Value of Exempt Supplies", "Value of Out-of-Scope Supplies",
            "Your Taxable Activity No."]
    rows = [head] + [[r["customer_tin"], r["customer"], r["invoice_no"],
                      r["date"], r["standard"] or None, r["zero"] or None,
                      r["exempt"] or None, r["oos"] or None, r["activity_no"]]
                     for r in data["output"]]
    t = data["output_total"]
    summary = [["Your Taxable Activity No."] + head[4:8],
               [data["company"]["activity_no"], t["standard"], t["zero"],
                t["exempt"], t["oos"]]]
    return _book([("Output Tax Statement", rows,
                   [16, 34, 18, 13, 26, 18, 18, 18, 16], [1]),
                  ("Summary", summary, [18, 26, 18, 18, 18], [1])])


def input_statement_xlsx(data):
    """MIRA's Input Tax Statement v25.1."""
    co = data["company"]
    head = ["#", "Supplier TIN", "Supplier Name", "Supplier Invoice Number",
            "Invoice Date", "Invoice Total (excluding GST)"] + [
        f"GST Charged at {r}%" for r in INPUT_RATES] + [
        "Your Taxable Activity Number", "Revenue / Capital"]
    rows = [["Input Tax Statement"], [],
            ["TIN:", co["tin"]], ["Taxpayer Name:", co["name"]],
            ["Taxable Period:", f"{data['date_from']:%d/%m/%Y} to "
                                f"{data['date_to']:%d/%m/%Y}"], [], head]
    for i, r in enumerate(data["input"], 1):
        rows.append([i, r["supplier_tin"], r["supplier"], r["invoice_no"],
                     r["date"], r["net"]]
                    + [r["gst"][str(x)] or None for x in INPUT_RATES]
                    + [r["activity_no"], r["kind"]])
    totals = ["Totals", "", "", "", "", data["input_total"]["net"]] + [
        sum((r["gst"][str(x)] for r in data["input"]), ZERO)
        for x in INPUT_RATES]
    rows += [totals, [], ["", "", "", "", "", "", "", "Total Input Tax", "",
                          "", data["input_total"]["gst"]]]
    return _book([("Input Tax Statement", rows,
                   [5, 16, 32, 20, 13, 18, 12, 12, 12, 12, 12, 16, 14],
                   [1, 7, len(rows) - 2])])
