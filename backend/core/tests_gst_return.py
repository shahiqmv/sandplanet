"""The GST return's figures and MIRA's two statements, out of the books
(MIRA 205 v25.1, Input / Output Tax Statement v25.1)."""
import io
from datetime import date
from decimal import Decimal as D

from openpyxl import load_workbook
from rest_framework.test import APIClient

from .tests_books_credit import CreditBase

URL = "/api/v1/ledger/reports/gst?from=2026-03-01&to=2026-03-31"


class GstReturnTests(CreditBase):
    def setUp(self):
        super().setUp()
        ok = self.assertEqual
        # sales: a standard-rated invoice, a zero-rated dollar invoice, a
        # cash sale; and money in that is not a sale at all
        ok(self.invoice(party_tin="1055555GST501").status_code, 201)
        ok(self.invoice(party="Kuramathi", reference="INV-2026-0008",
                        currency="USD", fx_rate="15.42", lines=[
                            {"account": self.acc["4120"].id, "amount": "2000",
                             "gst_treatment": "ZERO"}]).status_code, 201)
        ok(self.txn("DEPOSIT", party="Walk-in", lines=[
            {"account": self.acc["4120"].id, "amount": "1000",
             "gst_treatment": "STANDARD"}]).status_code, 201)
        ok(self.txn("DEPOSIT", party="Director", lines=[
            {"account": self.acc["2110"].id, "amount": "500000"}])
           .status_code, 201)
        # purchases: with a tax invoice, without one, and a capital item
        ok(self.txn("EXPENSE", party="STELCO", party_tin="1000123GST501",
                    tax_invoice_held=True, tax_invoice_no="ST-551", lines=[
                        {"account": self.acc["6320"].id, "amount": "1000",
                         "gst_treatment": "STANDARD"}]).status_code, 201)
        ok(self.txn("EXPENSE", party="Corner Shop", lines=[
            {"account": self.acc["6320"].id, "amount": "500",
             "gst_treatment": "STANDARD"}]).status_code, 201)
        ok(self.bill(tax_invoice_held=True, party_tin="1012345GST501",
                     lines=[{"account": self.acc["1710"].id,
                             "amount": "12000", "gst_treatment": "STANDARD"}])
           .status_code, 201)
        # not this period, voided, and from before the books
        ok(self.invoice(reference="INV-2026-0020", date="2026-04-02")
           .status_code, 201)
        v = self.invoice(reference="INV-2026-0009").data["id"]
        ok(self.c.post(f"/api/v1/ledger/txns/{v}/void", {"reason": "x"},
                       format="json").status_code, 200)
        ok(self.invoice(reference="INV-2025-0412", date="2025-11-30",
                        is_opening=True, amount="8000", lines=[])
           .status_code, 201)

    def test_the_boxes_of_the_return(self):
        r = self.c.get(URL)
        self.assertEqual(r.status_code, 200, r.data)
        d = r.data
        self.assertEqual(d["form"], "MIRA 205 v25.1")
        box = {b["box"]: b["amount"] for b in d["boxes"]}
        self.assertEqual(box, {
            1: D("55080"),      # 50,000 + 4,000 and 1,000 + 80, GST included
            2: D("30840"),      # USD 2,000 at 15.42
            3: D("0"), 4: D("0"), 5: D("85920"),
            6: D("4080"), 7: D("1040"), 8: D("0"), 9: D("0"),
            10: D("3040")})
        self.assertIn("at 8%", d["boxes"][0]["label"])
        # boxes 8 and 9 are the accountant's to fill
        self.assertEqual([b["box"] for b in d["boxes"] if not b["derived"]],
                         [8, 9])
        self.assertEqual(d["exact"], {"output_tax": D("4080.00"),
                                      "input_tax": D("1040.00"),
                                      "net": D("3040.00")})

    def test_the_two_statements(self):
        d = self.c.get(URL).data
        self.assertEqual(
            [(r["customer"], r["invoice_no"], r["standard"], r["zero"])
             for r in d["output"]],
            [("Soneva Fushi", "INV-2026-0007", D("50000.00"), D("0")),
             ("Kuramathi", "INV-2026-0008", D("0"), D("30840.00")),
             ("Walk-in", "DEP-001", D("1000.00"), D("0"))])
        self.assertEqual(d["output"][0]["customer_tin"], "1055555GST501")
        # only purchases with a tax invoice held; the fixed asset is capital
        self.assertEqual(
            [(r["supplier"], r["invoice_no"], r["net"], r["gst"]["8"],
              r["kind"]) for r in d["input"]],
            [("STELCO", "ST-551", D("1000.00"), D("80.00"), "Revenue"),
             ("Manas Hardware", "MH-1001", D("12000.00"), D("960.00"),
              "Capital")])
        self.assertEqual(d["input_total"], {"net": D("13000.00"),
                                            "gst": D("1040.00")})

    def test_what_else_touched_the_gst_accounts(self):
        # paying MIRA, and a journal straight to the GST account
        self.txn("EXPENSE", party="MIRA", date="2026-03-28", lines=[
            {"account": self.acc["2210"].id, "amount": "3000"}])
        self.journal([self.line("2210", credit=50), self.line("6320", debit=50)])
        d = self.c.get(URL).data
        self.assertEqual(sorted((o["debit"], o["credit"])
                                for o in d["other_entries"]),
                         [(D("0.00"), D("50.00")), (D("3000.00"), D("0.00"))])
        # the return itself is unmoved by them
        self.assertEqual(d["boxes"][5]["amount"], D("4080"))

    def test_warnings_for_what_mira_would_send_back(self):
        self.invoice(reference="INV-2026-0030", lines=[
            {"account": self.acc["4120"].id, "amount": "700"}])
        # the form refuses a claim without the TIN and the invoice number …
        body = {"party": "No TIN Traders", "tax_invoice_held": True,
                "tax_invoice_date": "2024-12-01", "lines": [
                    {"account": self.acc["6320"].id, "amount": "100",
                     "gst_treatment": "STANDARD"}]}
        self.assertIn("To claim the GST",
                      self.txn("EXPENSE", **body).data["detail"])
        # … an entry from before that rule still shows up as needing them
        r = self.txn("EXPENSE", party_tin="1", tax_invoice_no="X", **body)
        self.assertEqual(r.status_code, 201, r.data)
        from .models import LedgerTxn
        LedgerTxn.objects.filter(pk=r.data["id"]).update(party_tin="")
        w = " ".join(self.c.get(URL).data["warnings"])
        self.assertIn("1 invoice line has no GST treatment", w)
        self.assertIn("lacks the supplier's TIN", w)
        self.assertIn("more than 12 months", w)
        self.assertIn("taxable activity number", w)

    def test_the_statements_as_mira_lays_them_out(self):
        r = self.c.get(URL + "&export=output")
        self.assertIn("spreadsheet", r["Content-Type"])
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb["Output Tax Statement"]
        self.assertEqual([c.value for c in ws[1]], [
            "Customer TIN", "Customer Name", "Invoice No.", "Invoice Date",
            "Value of Supplies Subject to GST at 8% or 17% (excluding GST)",
            "Value of Zero-Rated Supplies", "Value of Exempt Supplies",
            "Value of Out-of-Scope Supplies", "Your Taxable Activity No."])
        self.assertEqual((ws["B2"].value, ws["E2"].value), ("Soneva Fushi",
                                                           50000))
        self.assertEqual(wb["Summary"]["B2"].value, 51000)
        r = self.c.get(URL + "&export=input")
        ws = load_workbook(io.BytesIO(r.content))["Input Tax Statement"]
        self.assertEqual([c.value for c in ws[7]], [
            "#", "Supplier TIN", "Supplier Name", "Supplier Invoice Number",
            "Invoice Date", "Invoice Total (excluding GST)",
            "GST Charged at 6%", "GST Charged at 8%", "GST Charged at 12%",
            "GST Charged at 16%", "GST Charged at 17%",
            "Your Taxable Activity Number", "Revenue / Capital"])
        self.assertEqual((ws["C8"].value, ws["H8"].value, ws["M9"].value),
                         ("STELCO", 80, "Capital"))
        self.assertEqual(ws["A6"].value, None)
        self.assertIn("01/03/2026 to 31/03/2026", ws["B5"].value)

    def test_a_period_is_needed_and_the_books_roles_only(self):
        self.assertEqual(self.c.get("/api/v1/ledger/reports/gst").status_code,
                         400)
        pm = APIClient()
        pm.force_authenticate(self.pm)
        self.assertEqual(pm.get(URL).status_code, 403)
        self.assertIsInstance(self.c.get(URL).data["date_to"], date)
