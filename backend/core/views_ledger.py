"""The books — API (`/api/v1/ledger/*`). Rules live in core/ledger.py."""
from io import BytesIO

from django.http import HttpResponse
from rest_framework.decorators import api_view
from rest_framework.response import Response

from . import ledger
from .models import JournalEntry, LedgerAccount


def _read(request):
    if request.user.role not in ledger.READ_ROLES:
        return Response({"detail": "The books are open to Finance, the "
                                   "Director and the signatory."}, status=403)
    return None


def _write(request):
    if request.user.role not in ledger.WRITE_ROLES:
        return Response({"detail": "Finance keeps the books."}, status=403)
    return None


def _account(a, balances=None):
    bal = (balances or {}).get(a.id)
    return {
        "id": a.id, "code": a.code, "name": a.name, "type": a.type,
        "type_label": a.get_type_display(), "parent": a.parent_id,
        "is_group": a.is_group, "currency": a.currency,
        "system": bool(a.system_key), "is_bank": bool(a.bank_account_id),
        "description": a.description, "is_active": a.is_active,
        "debit_normal": a.debit_normal,
        # on the account's normal side: a liability shows its credit balance
        "balance": None if bal is None else (bal if a.debit_normal else -bal),
    }


def _entry(e, lines=True):
    out = {
        "id": e.id, "ref": e.ref, "date": e.date, "kind": e.kind,
        "kind_label": e.get_kind_display(), "status": e.status, "memo": e.memo,
        "source_ref": e.source_ref,
        "reversal_of": e.reversal_of.ref if e.reversal_of_id else None,
        "reversed_by": (e.reversed_by.ref
                        if getattr(e, "reversed_by", None) else None),
        "created_by": e.created_by.full_name if e.created_by_id else "",
        "posted_by": e.posted_by.full_name if e.posted_by_id else "",
        "posted_at": e.posted_at,
    }
    rows = list(e.lines.select_related("account", "site", "project",
                                       "cost_head"))
    out["total"] = sum((ln.debit for ln in rows), ledger.ZERO)
    if lines:
        out["lines"] = [{
            "id": ln.id, "account": ln.account_id,
            "account_code": ln.account.code, "account_name": ln.account.name,
            "description": ln.description, "debit": ln.debit,
            "credit": ln.credit, "currency": ln.currency,
            "amount_fc": ln.amount_fc, "fx_rate": ln.fx_rate,
            "site": ln.site_id, "site_code": ln.site.code if ln.site_id else "",
            "project": ln.project_id,
            "project_code": ln.project.code if ln.project_id else "",
            "cost_head": ln.cost_head_id, "party": ln.party,
        } for ln in rows]
    return out


# ---- chart of accounts ------------------------------------------------------

@api_view(["GET", "POST"])
def accounts(request):
    if (bad := _read(request)):
        return bad
    if request.method == "POST":
        if (bad := _write(request)):
            return bad
        a, msg = ledger.save_account(request.data, request.user)
        if msg:
            return Response({"detail": msg}, status=400)
        return Response(_account(a), status=201)
    from django.db.models import Sum
    from .models import JournalLine
    bal = {r["account"]: (r["d"] or 0) - (r["c"] or 0)
           for r in JournalLine.objects.filter(entry__status="POSTED")
           .values("account").annotate(d=Sum("debit"), c=Sum("credit"))}
    rows = list(LedgerAccount.objects.all())
    return Response({
        "accounts": [_account(a, bal) for a in rows],
        "types": [{"value": v, "label": lab}
                  for v, lab in LedgerAccount.Type.choices],
        "can_edit": request.user.role in ledger.WRITE_ROLES,
        "settings": ledger.settings_dict(),
        "posted_entries": JournalEntry.objects.filter(status="POSTED").count(),
    })


@api_view(["PATCH", "DELETE"])
def account_detail(request, pk):
    if (bad := _write(request)):
        return bad
    a = LedgerAccount.objects.filter(pk=pk).first()
    if a is None:
        return Response({"detail": "Not found."}, status=404)
    if request.method == "DELETE":
        msg = ledger.delete_account(a, request.user)
        if msg:
            return Response({"detail": msg}, status=400)
        return Response(status=204)
    a, msg = ledger.save_account(request.data, request.user, account=a)
    if msg:
        return Response({"detail": msg}, status=400)
    return Response(_account(a))


@api_view(["POST"])
def setup(request):
    """Lay down the standard chart on empty books."""
    if (bad := _write(request)):
        return bad
    made, msg = ledger.setup_standard_chart(request.user)
    if msg:
        return Response({"detail": msg}, status=400)
    return Response({"accounts": made}, status=201)


@api_view(["POST"])
def sync_banks(request):
    if (bad := _write(request)):
        return bad
    return Response({"added": ledger.sync_bank_accounts(request.user)})


@api_view(["GET", "POST"])
def settings(request):
    if (bad := _read(request)):
        return bad
    if request.method == "POST":
        if request.user.role != "ADMIN" and "books_start_date" in request.data:
            return Response({"detail": "Only an admin moves the start of the "
                                       "books."}, status=403)
        if (bad := _write(request)):
            return bad
        msg = ledger.save_settings(request.data, request.user)
        if msg:
            return Response({"detail": msg}, status=400)
    return Response(ledger.settings_dict())


# ---- journals ---------------------------------------------------------------

@api_view(["GET", "POST"])
def journals(request):
    if (bad := _read(request)):
        return bad
    if request.method == "POST":
        if (bad := _write(request)):
            return bad
        e, msg = ledger.save_draft(request.data, request.user)
        if msg:
            return Response({"detail": msg}, status=400)
        if request.data.get("post"):
            msg = ledger.post(e, request.user)
            if msg:                      # kept as a draft, with the reason
                return Response({"detail": msg, "draft": _entry(e)},
                                status=400)
            e.refresh_from_db()
        return Response(_entry(e), status=201)
    qs = ledger.journal_filter(
        JournalEntry.objects.select_related("created_by", "posted_by",
                                            "reversal_of", "reversed_by"),
        request.GET)
    try:
        limit = min(int(request.GET.get("limit", 50)), 200)
        offset = max(int(request.GET.get("offset", 0)), 0)
    except ValueError:
        limit, offset = 50, 0
    total = qs.count()
    return Response({
        "journals": [_entry(e, lines=False) for e in qs[offset:offset + limit]],
        "total": total, "has_more": offset + limit < total,
        "drafts": JournalEntry.objects.filter(status="DRAFT").count(),
        "can_edit": request.user.role in ledger.WRITE_ROLES,
    })


def _get_entry(pk):
    return (JournalEntry.objects.select_related(
        "created_by", "posted_by", "reversal_of", "reversed_by")
        .filter(pk=pk).first())


@api_view(["GET", "PATCH", "DELETE"])
def journal_detail(request, pk):
    if (bad := _read(request)):
        return bad
    e = _get_entry(pk)
    if e is None:
        return Response({"detail": "Not found."}, status=404)
    if request.method == "GET":
        return Response(_entry(e))
    if (bad := _write(request)):
        return bad
    if request.method == "DELETE":
        if e.status != "DRAFT":
            return Response({"detail": "A posted entry is never deleted — "
                                       "reverse it."}, status=400)
        e.delete()
        return Response(status=204)
    e, msg = ledger.save_draft(request.data, request.user, entry=e)
    if msg:
        return Response({"detail": msg}, status=400)
    if request.data.get("post"):
        msg = ledger.post(e, request.user)
        if msg:
            return Response({"detail": msg, "draft": _entry(e)}, status=400)
    return Response(_entry(_get_entry(pk)))


@api_view(["POST"])
def journal_action(request, pk, action):
    if (bad := _write(request)):
        return bad
    e = _get_entry(pk)
    if e is None:
        return Response({"detail": "Not found."}, status=404)
    if action == "post":
        msg = ledger.post(e, request.user)
        if msg:
            return Response({"detail": msg}, status=400)
        return Response(_entry(_get_entry(pk)))
    if action == "reverse":
        rev, msg = ledger.reverse(
            e, request.user, on=ledger._as_date(request.data.get("date")),
            reason=request.data.get("reason"))
        if msg:
            return Response({"detail": msg}, status=400)
        return Response(_entry(_get_entry(rev.id)), status=201)
    return Response({"detail": "Unknown action."}, status=400)


# ---- reports ----------------------------------------------------------------

def _xlsx(title, subtitle, headers, rows, totals=None, widths=None,
          name="report"):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = title[:28]
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = subtitle
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True)
    r = 5
    for row in rows:
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            if isinstance(v, (int, float)) or hasattr(v, "quantize"):
                c.number_format = "#,##0.00"
        r += 1
    if totals:
        for i, v in enumerate(totals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(bold=True)
            if hasattr(v, "quantize"):
                c.number_format = "#,##0.00"
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(
        buf.getvalue(), content_type="application/vnd.openxmlformats-"
        "officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
    return resp


@api_view(["GET"])
def trial_balance(request):
    if (bad := _read(request)):
        return bad
    tb = ledger.trial_balance(ledger._as_date(request.GET.get("from")),
                              ledger._as_date(request.GET.get("to")))
    if request.GET.get("export") == "xlsx":
        t = tb["totals"]
        return _xlsx(
            "Trial balance",
            f"{tb['date_from']:%d %b %Y} to {tb['date_to']:%d %b %Y} · MVR",
            ["Code", "Account", "Type", "Opening Dr", "Opening Cr",
             "Debit", "Credit", "Closing Dr", "Closing Cr"],
            [[r["code"], r["name"], r["type_label"], r["opening_debit"],
              r["opening_credit"], r["debit"], r["credit"],
              r["closing_debit"], r["closing_credit"]] for r in tb["rows"]],
            ["", "Total", "", t["opening_debit"], t["opening_credit"],
             t["debit"], t["credit"], t["closing_debit"],
             t["closing_credit"]],
            [9, 44, 14, 15, 15, 15, 15, 15, 15],
            f"trial-balance-{tb['date_to']}")
    return Response(tb)


@api_view(["GET"])
def account_ledger(request, pk):
    if (bad := _read(request)):
        return bad
    a = LedgerAccount.objects.filter(pk=pk).first()
    if a is None:
        return Response({"detail": "Not found."}, status=404)
    if a.is_group:
        return Response({"detail": "A group holds no entries of its own — "
                                   "open an account under it."}, status=400)
    led = ledger.account_ledger(a, ledger._as_date(request.GET.get("from")),
                                ledger._as_date(request.GET.get("to")))
    if request.GET.get("export") == "xlsx":
        return _xlsx(
            f"{a.code} {a.name}",
            f"{led['date_from']:%d %b %Y} to {led['date_to']:%d %b %Y} · MVR",
            ["Date", "Entry", "Memo", "Description", "Party", "Site",
             "Debit", "Credit", "Balance"],
            [["", "", "Brought forward", "", "", "", "", "", led["opening"]]]
            + [[r["date"], r["ref"], r["memo"], r["description"], r["party"],
                r["site"], r["debit"] or None, r["credit"] or None,
                r["balance"]] for r in led["rows"]],
            ["", "", "Total", "", "", "", led["debit"], led["credit"],
             led["closing"]],
            [12, 10, 40, 34, 24, 8, 15, 15, 16],
            f"ledger-{a.code}-{led['date_to']}")
    return Response(led)
