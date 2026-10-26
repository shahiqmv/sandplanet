"""Fill in the language a worker's forms print in (owner 2026-10-26).

From nationality where it settles it — Bangladeshi → Bengali, Nepali →
Nepali, Maldivian → Dhivehi (the English form) — and, with --sheet, from the
reviewed spreadsheet for the Sri Lankans and Indians (column A emp no,
F the suggested language, I the confirmed one, which wins).

Only a blank language is filled; a language already on a record is left
alone. A dry run unless --apply.

    python manage.py set_worker_languages
    python manage.py set_worker_languages --sheet /tmp/langs.xlsx --apply
"""
from collections import Counter

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from core.audit import audit
from core.models import Employee

NAMES = {"bengali": "BN", "tamil": "TA", "sinhala": "SI", "hindi": "HI",
         "nepali": "NE", "dhivehi": "DV", "english": "EN"}


def from_nationality(nat):
    n = (nat or "").strip().lower().replace("'", "")
    if n.startswith("ban") or "desh" in n or n.startswith("bdesh"):
        return "BN"
    if n.startswith("nepal"):
        return "NE"
    if n.startswith("maldiv") or n == "male":
        return "DV"
    return None


def from_sheet(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        word = (row[8] if len(row) > 8 and row[8] else row[5]) or ""
        code = NAMES.get(str(word).strip().split(" ")[0].lower())
        if code:
            out[str(row[0]).strip()] = code
    return out


class Command(BaseCommand):
    help = "Fill blank worker languages from nationality and a reviewed sheet."

    def add_arguments(self, parser):
        parser.add_argument("--sheet", help="the reviewed .xlsx")
        parser.add_argument("--apply", action="store_true")

    def handle(self, *args, **opts):
        by_sheet = {}
        if opts.get("sheet"):
            try:
                by_sheet = from_sheet(opts["sheet"])
            except Exception as e:   # noqa: BLE001 — a bad file is the user's to fix
                raise CommandError(f"Could not read the sheet: {e}") from e
        plan, counts, left = [], Counter(), Counter()
        for e in Employee.objects.filter(preferred_language="").only(
                "id", "emp_no", "nationality"):
            code = by_sheet.get(e.emp_no) or from_nationality(e.nationality)
            if code:
                plan.append((e.id, code))
                counts[code] += 1
            else:
                left[(e.nationality or "—").strip()] += 1
        self.stdout.write(f"To set: {sum(counts.values())} "
                          f"{dict(counts)}")
        self.stdout.write(f"Still blank after this: {sum(left.values())} "
                          f"{dict(left.most_common(10))}")
        if not opts["apply"]:
            self.stdout.write("Dry run — nothing written. Add --apply.")
            return
        with transaction.atomic():
            for pk, code in plan:
                Employee.objects.filter(pk=pk, preferred_language="").update(
                    preferred_language=code)
        audit("employee", 0, "LANGUAGES_FILLED", actor=None,
              detail={"set": dict(counts), "sheet": bool(by_sheet)})
        self.stdout.write(self.style.SUCCESS(
            f"Set {sum(counts.values())} languages."))
