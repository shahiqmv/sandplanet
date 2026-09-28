"""One person's cost over a date range — attendance, overtime and what the
payroll runs paid them (owner 2026-09-28: "report of individual staff cost
date range. like OT hrs, attendance, salary summary").

Attendance and overtime are read day by day inside the range. Pay comes from
the payroll lines of every run whose month touches the range — a run is a
whole month, so a range that starts or ends mid-month shows that month's full
line and says so rather than inventing a pro-rata figure payroll never paid.
"""
from collections import OrderedDict
from datetime import date
from decimal import Decimal

from .models import Attendance, PayrollLine, SalaryAdvance, Site
from .payroll import compute_line, friday_ot_hours, month_days

MAX_DAYS = 400
ZERO = Decimal("0")
MARKS = OrderedDict([("PRESENT", "Present"), ("HALF_DAY", "Half day"),
                     ("ABSENT", "Absent"), ("SICK", "Sick"),
                     ("PAID_LEAVE", "Leave (paid)"),
                     ("LEAVE", "Leave (no pay)")])
MONTHS = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep",
          "Oct", "Nov", "Dec"]


def _months(start, end):
    y, m = start.year, start.month
    while (y, m) <= (end.year, end.month):
        yield y, m
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)


def build(emp, start, end):
    if end < start:
        raise ValueError("The end date is before the start date.")
    if (end - start).days + 1 > MAX_DAYS:
        raise ValueError(f"Pick a range of {MAX_DAYS} days or less.")
    rate, rate_ccy = emp.ot_terms()
    sub = emp.engagement_type == "SUBCONTRACT"
    weeks = {s.id: set(s.working_days or [6, 7, 1, 2, 3, 4])
             for s in Site.objects.all()}

    # ---- attendance, day by day -------------------------------------------
    rows = list(Attendance.objects.filter(employee=emp, day__gte=start,
                                          day__lte=end)
                .select_related("site").order_by("day"))
    counts = OrderedDict((k, 0) for k in MARKS)
    by_site = OrderedDict()
    ot = {"requested": ZERO, "approved": ZERO, "pending": ZERO,
          "pending_days": 0, "workday": ZERO, "rest_day": ZERO}
    rest_worked = 0
    days = []
    for a in rows:
        rest = a.day.isoweekday() not in weeks.get(a.site_id, set())
        counts[a.remark] = counts.get(a.remark, 0) + 1
        s = by_site.setdefault(a.site.code, {"site": a.site.code,
                                             "name": a.site.name, "days": 0,
                                             "worked": Decimal("0"),
                                             "ot": ZERO})
        s["days"] += 1
        if a.remark == "PRESENT":
            s["worked"] += 1
        elif a.remark == "HALF_DAY":
            s["worked"] += Decimal("0.5")
        req = (a.sub_extra_hours if sub else a.ot_requested) or ZERO
        appr = a.sub_extra_approved if sub else a.ot_approved
        ot["requested"] += req
        if appr is None:
            if req > 0:
                ot["pending"] += req
                ot["pending_days"] += 1
        else:
            ot["approved"] += appr
            s["ot"] += appr
            # A worked rest day is paid as the flat Friday allowance, never
            # its hours as well (payroll._attendance_prefill).
            if rest and a.remark == "PRESENT":
                ot["rest_day"] += appr
            else:
                ot["workday"] += appr
        if rest and a.remark == "PRESENT":
            rest_worked += 1
        days.append({
            "day": a.day, "dow": a.day.strftime("%a"), "site": a.site.code,
            "rest": rest, "remark": a.remark,
            "mark": MARKS.get(a.remark, a.remark),
            "in": a.check_in.strftime("%H:%M") if a.check_in else "",
            "out": a.check_out.strftime("%H:%M") if a.check_out else "",
            "ot_requested": req, "ot_approved": appr,
            "pending": appr is None and req > 0,
        })
    # Days inside his employment the register never mentions
    first = max(start, emp.join_date) if emp.join_date else start
    last = min(end, emp.left_on) if emp.left_on else end
    employed_days = max((last - first).days + 1, 0)
    unmarked = max(employed_days - len(rows), 0)
    worked = counts["PRESENT"] + Decimal(counts["HALF_DAY"]) / 2

    # ---- pay: the runs whose month touches the range ----------------------
    months = list(_months(start, end))
    lines = list(PayrollLine.objects.filter(employee=emp).filter(
        run__year__gte=start.year, run__year__lte=end.year)
        .select_related("run", "run__site", "site")
        .order_by("run__year", "run__month", "run__currency", "run__id"))
    lines = [ln for ln in lines if (ln.run.year, ln.run.month) in months]
    fri_h = friday_ot_hours()
    pay, totals, seen = [], {}, set()
    for ln in lines:
        run = ln.run
        money = compute_line(ln, fri_h)
        seen.add((run.year, run.month))
        m_start = date(run.year, run.month, 1)
        m_end = date(run.year, run.month, month_days(run.year, run.month))
        pay.append({
            "run": run.ref or f"{run.year}-{run.month:02d}",
            "run_id": run.id, "line_id": ln.id,
            "period": f"{MONTHS[run.month]} {run.year}",
            "kind": run.get_kind_display(),
            "site": run.site.code if run.site_id else "All sites",
            "currency": run.currency, "status": run.get_status_display(),
            "locked": run.status == "LOCKED", "excluded": ln.excluded,
            "part_month": m_start < start or m_end > end,
            "days_worked": ln.days_worked, "fridays_worked": ln.fridays_worked,
            "ot_hours": ln.ot_hours, "ot_rate": ln.ot_rate,
            "basic_pay": ln.basic_pay, **money,
            "advance": ln.advance, "loan": ln.loan, "penalty": ln.penalty,
        })
        t = totals.setdefault(run.currency, {
            "currency": run.currency, "earned_basic": ZERO, "friday_pay": ZERO,
            "ot_pay": ZERO, "allowance": ZERO, "gross": ZERO,
            "deductions": ZERO, "net": ZERO, "ot_hours": ZERO, "fridays": 0})
        for k in ("earned_basic", "friday_pay", "ot_pay", "allowance",
                  "gross", "deductions", "net"):
            t[k] += money[k]
        t["ot_hours"] += ln.ot_hours
        t["fridays"] += ln.fridays_worked
    # Only months he was employed in: a month before he joined has no run
    # to be missing from.
    j = (emp.join_date.year, emp.join_date.month) if emp.join_date else None
    lv = (emp.left_on.year, emp.left_on.month) if emp.left_on else None
    not_run = [f"{MONTHS[m]} {y}" for (y, m) in months
               if (y, m) not in seen and (not j or (y, m) >= j)
               and (not lv or (y, m) <= lv)]

    advances = [{
        "ref": a.document.ref, "kind": a.get_kind_display(),
        "amount": a.amount, "months": a.months,
        "from": f"{MONTHS[a.period_month]} {a.period_year}",
        "status": (a.document.status or "").replace("_", " ").capitalize(),
    } for a in SalaryAdvance.objects.filter(
        employee=emp, created_at__date__gte=start, created_at__date__lte=end)
        .exclude(document__status__in=("CANCELLED", "VOID", "REJECTED",
                                       "WITHDRAWN"))
        .select_related("document").order_by("created_at")]

    return {
        "employee": {
            "id": emp.id, "emp_no": emp.emp_no, "full_name": emp.full_name,
            "category": emp.job_category.name if emp.job_category_id else "",
            "job_title": emp.job_title or "",
            "currency": emp.currency or "MVR", "basic_pay": emp.basic_pay,
            "usd_basic_pay": emp.usd_basic_pay,
            "ot_rate": rate, "ot_currency": rate_ccy,
            "employment_type": emp.get_employment_type_display(),
            "subcontract": sub,
            "subcontractor": (emp.subcontractor.name
                              if sub and emp.subcontractor_id else ""),
            "join_date": emp.join_date, "left_on": emp.left_on,
            "is_active": emp.is_active,
        },
        "start": start, "end": end,
        "calendar_days": (end - start).days + 1,
        "employed_days": employed_days,
        "attendance": {
            "counts": [{"key": k, "label": MARKS[k], "days": counts[k]}
                       for k in MARKS],
            "marked": len(rows), "unmarked": unmarked, "worked": worked,
            "rest_days_worked": rest_worked, "by_site": list(by_site.values()),
        },
        "ot": {**ot, "rate": rate, "currency": rate_ccy,
               "value_workday": (ot["workday"] * rate).quantize(Decimal("0.01")),
               "friday_hours": fri_h},
        "pay": pay, "pay_totals": list(totals.values()), "not_run": not_run,
        "any_part_month": any(p["part_month"] for p in pay),
        "advances": advances, "days": days,
    }


def workbook(rep):
    """The same report as a workbook: summary, pay by month, every day."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F3A5F")
    bold = Font(bold=True)
    e = rep["employee"]

    def header(ws, row, cols):
        for i, c in enumerate(cols, 1):
            x = ws.cell(row=row, column=i, value=c)
            x.font, x.fill = head, fill
            x.alignment = Alignment(wrap_text=True, vertical="center")

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Staff cost — {e['emp_no']} {e['full_name']}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{rep['start']:%d %b %Y} to {rep['end']:%d %b %Y}"
    info = [("Category", e["category"]), ("Job title", e["job_title"]),
            ("Employment", e["employment_type"]),
            ("Salary currency", e["currency"]),
            ("Basic pay (monthly)", e["basic_pay"]),
            ("USD basic", e["usd_basic_pay"] or None),
            ("OT rate", f"{e['ot_currency']} {e['ot_rate']}/h"
             if e["ot_rate"] else "none"),
            ("Joined", e["join_date"]), ("Left", e["left_on"]), ("", ""),
            ("Days in range", rep["calendar_days"]),
            ("Days employed in range", rep["employed_days"])]
    info += [(c["label"], c["days"]) for c in rep["attendance"]["counts"]]
    info += [("Days worked (half day = ½)", rep["attendance"]["worked"]),
             ("Rest days worked", rep["attendance"]["rest_days_worked"]),
             ("Days not on the register", rep["attendance"]["unmarked"]),
             ("", ""),
             ("OT hours requested", rep["ot"]["requested"]),
             ("OT hours approved", rep["ot"]["approved"]),
             ("…on working days (paid as OT)", rep["ot"]["workday"]),
             ("…on worked rest days (covered by the flat Friday pay)",
              rep["ot"]["rest_day"]),
             ("OT hours awaiting approval", rep["ot"]["pending"])]
    for i, (k, v) in enumerate(info, 4):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v).font = bold
    r = 4 + len(info) + 1
    if rep["pay_totals"]:
        ws.cell(row=r, column=1, value="Paid by payroll in these months").font = bold
        header(ws, r + 1, ["Currency", "Earned basic", "Friday pay", "OT pay",
                           "Allowance", "Gross", "Deductions", "Net"])
        for j, t in enumerate(rep["pay_totals"], r + 2):
            for c, k in enumerate(["currency", "earned_basic", "friday_pay",
                                   "ot_pay", "allowance", "gross",
                                   "deductions", "net"], 1):
                ws.cell(row=j, column=c, value=t[k])
    widths(ws, [44, 16, 12, 12, 12, 12, 12, 12])

    ws2 = wb.create_sheet("Pay by month")
    cols = ["Run", "Period", "Site", "Currency", "Status", "Days worked",
            "Fridays", "OT h", "OT rate", "Basic (monthly)", "Earned basic",
            "Friday pay", "OT pay", "Allowance", "Gross", "Advance", "Loan",
            "Penalty", "Net", "Note"]
    header(ws2, 1, cols)
    for i, p in enumerate(rep["pay"], 2):
        note = "; ".join(x for x in (
            "whole month — range covers part" if p["part_month"] else "",
            "excluded (settled in cash)" if p["excluded"] else "",
            "" if p["locked"] else "not locked yet") if x)
        vals = [p["run"], p["period"], p["site"], p["currency"], p["status"],
                p["days_worked"], p["fridays_worked"], p["ot_hours"],
                p["ot_rate"], p["basic_pay"], p["earned_basic"],
                p["friday_pay"], p["ot_pay"], p["allowance"], p["gross"],
                p["advance"], p["loan"], p["penalty"], p["net"], note]
        for j, v in enumerate(vals, 1):
            ws2.cell(row=i, column=j, value=v)
    if rep["not_run"]:
        ws2.cell(row=len(rep["pay"]) + 3, column=1,
                 value="No payroll run yet: " + ", ".join(rep["not_run"]))
    widths(ws2, [14, 10, 9, 8, 14] + [9] * 14 + [30])

    ws3 = wb.create_sheet("Daily")
    header(ws3, 1, ["Date", "Day", "Site", "Rest day", "Mark", "In", "Out",
                    "OT requested", "OT approved", "Awaiting approval"])
    for i, d in enumerate(rep["days"], 2):
        vals = [d["day"], d["dow"], d["site"], "yes" if d["rest"] else "",
                d["mark"], d["in"], d["out"], d["ot_requested"],
                d["ot_approved"], "yes" if d["pending"] else ""]
        for j, v in enumerate(vals, 1):
            c = ws3.cell(row=i, column=j, value=v)
            if j == 1:
                c.number_format = "dd mmm yyyy"
    ws3.freeze_panes = "A2"
    widths(ws3, [12, 6, 7, 8, 14, 7, 7, 11, 11, 10])

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
